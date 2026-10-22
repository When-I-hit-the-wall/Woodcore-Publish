"""The project's positions, Woodcore's own (calc/positions_store.py, v0.114, docs/domain/positions.md §7): `Woodcore-Datei/positions.json` the one source of positions, Positionen.xlsx imported once and moved into the archives, and every field but the id and the type optional everywhere -- in the file, in the export, in Typst.

What is held here: the file read and written whole and atomically, a row refused only for what makes it no position; an id proposed with no gap ever refilled; a blank position (id and type only) made for the plan tab; a workbook's rows mapped, its gaps said in notes and never refused; the migration moving the workbook into `_Archiv/` and doing nothing the second time; the export writing a position whose section does not read as `section: null` with a warning; and a blank position going through `#member`, `#moisture-class-of`, `calc-block`, a schema and a connector without stopping anything.
"""

TIER = "fast"  # tests/run_tests.py --fast

import json
import shutil
import subprocess

import conftest
from conftest import Skip

import engine
import positions_store as store
import project
import validation


def _here(name):
    """A project of the test's own: a copy of the example's sources (its positions.json among them) and export."""
    return conftest.project_copy(name)


def _refused(code, call, *args):
    try:
        call(*args)
    except store.PositionsError as exc:
        assert exc.code == code, (code, exc.code, str(exc))
        return exc
    raise AssertionError(f"{code}: not refused")


def _workbook(path, rows, headers=("position", "type", "material", "resistance class", "section type", "dimension", "units", "moisture class", "description")):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "Positions Element"
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    book.save(path)
    return path


# --- the file -------------------------------------------------------------------------------

def test_the_file_is_read_and_written_whole_and_only_the_id_and_the_type_are_required():
    here = _here("store_file")
    rows = store.read(here)
    assert [row["position"] for row in rows] == ["T001", "S001", "S002"], rows
    assert set(rows[0]) == set(store.KEYS) and rows[0]["moisture_class"] == 2
    store.write(here, [{"position": "T001", "type": "Träger"}, {"position": "D001", "type": "Decke/Dach", "dimension": 180, "moisture_class": "2"}])
    raw = store.path(here).read_text(encoding="utf-8")
    assert "\r" not in raw and raw.endswith("\n") and "Träger" in raw, "LF, the umlaut written as it is"
    data = json.loads(raw)
    assert data["format"] == store.FORMAT and data["positions"][0] == {"position": "T001", "type": "Träger"}, "a blank key is not written"
    (t001, d001) = store.read(here)
    assert t001["material"] == "" and t001["moisture_class"] is None, "an absent key reads as blank"
    assert d001["dimension"] == "180" and d001["moisture_class"] == 2, "a number typed as one is a text, a moisture class as text a number"
    assert not list(store.path(here).parent.glob("*.tmp")), "the file is written beside itself and moved into place"
    _refused("no-type", store.write, here, [{"position": "T001"}])
    _refused("no-id", store.write, here, [{"type": "Träger"}])
    _refused("bad-id", store.write, here, [{"position": "T 001", "type": "Träger"}])
    _refused("duplicate-id", store.write, here, [{"position": "T001", "type": "Träger"}, {"position": "t001", "type": "Träger"}])
    _refused("moisture", store.write, here, [{"position": "T001", "type": "Träger", "moisture_class": 4}])
    assert [row["position"] for row in store.read(here)] == ["T001", "D001"], "a refused write leaves the file as it was"


def test_a_position_is_added_changed_renamed_and_deleted():
    here = _here("store_rows")
    store.add(here, {"position": "T002", "type": "Träger", "material": "Fichte"})
    _refused("duplicate-id", store.add, here, {"position": "t002", "type": "Träger"})
    row, renamed = store.update(here, "T002", {"position": "T002", "type": "Träger", "description": "neu"})
    assert renamed is None and row["description"] == "neu"
    row, renamed = store.update(here, "T002", {"position": "T009", "type": "Träger"})
    assert renamed == "T002" and [r["position"] for r in store.read(here)] == ["T001", "S001", "S002", "T009"], "renamed in its place"
    _refused("duplicate-id", store.update, here, "T009", {"position": "S001", "type": "Träger"})
    _refused("unknown-position", store.update, here, "T002", {"position": "T002", "type": "Träger"})
    store.delete(here, "T009")
    _refused("unknown-position", store.delete, here, "T009")
    assert [r["position"] for r in store.read(here)] == ["T001", "S001", "S002"]


def test_a_file_that_cannot_be_read_is_refused_and_never_read_as_no_positions():
    """Read as an empty list, a broken positions.json would have the next export sweep every data/positions/<id>.json."""
    here = _here("store_broken")
    store.path(here).write_text("{ not json", encoding="utf-8")
    _refused("unreadable", store.read, here)
    _refused("unreadable", store.add, here, {"position": "T009", "type": "Träger"})
    assert store.path(here).read_text(encoding="utf-8") == "{ not json", "a file that cannot be read is not written over"
    store.path(here).unlink()
    assert store.read(here) == [], "a project with no positions.json has no position, and that is no error"


def test_an_id_is_proposed_and_a_gap_is_never_refilled():
    assert store.propose_id(["T001", "T003"], "T") == "T004", "a number deleted may still be named somewhere: never taken again"
    assert store.propose_id(["T001"], "S") == "S001", "a symbol with no position starts on floor 0"
    assert store.propose_id(["T001"], "T", floor="1") == "T101"
    assert store.propose_id(["TU01", "TU02", "T001"], "T") == "T002", "the floor of the last position of that symbol"
    assert store.propose_id(["T001", "TU01"], "T") == "TU02"
    assert store.propose_id(["T099"], "T") == "T0100"
    assert store.propose_id(["t002"], "T", floor="0") == "T003", "Windows does not tell T003 from t003: taken ignoring case"
    assert store.propose_id([{"position": "S001"}, {"position": "S007"}], "S") == "S008"
    assert store.propose_id([], "S", floor="") == "S001"
    _refused("no-symbol", store.propose_id, [], "")


def test_the_numbers_count_up_by_prefix_and_storey_over_the_project_and_what_is_drawn_unsaved():
    """Issue #9 (REQ-003): S101, S102, S103 on the first floor, T101 for the first beam there, S201 on the second -- a counter per prefix and per storey, over the project's positions and the numbers drawn and not yet saved, given as ids beside the rows."""
    rows = [{"position": "S101", "type": "Stütze"}, {"position": "T201", "type": "Träger"}]
    drawn = ["S102"]
    assert store.propose_id(rows + drawn, "S", floor="1") == "S103"
    assert store.propose_id(rows + drawn, "T", floor="1") == "T101"
    assert store.propose_id(rows + drawn, "S", floor="2") == "S201"
    assert store.propose_id(rows + drawn, "T", floor="2") == "T202"


def test_the_section_proposed_is_the_last_of_the_same_type_copied_with_its_units():
    """Issue #9 (REQ-006, REQ-007): «la section proposée est celle du dernier objet du même type dessiné» -- a copy of its three cells, the units with them, so «10/20» in cm is not read as mm."""
    rows = [
        {"position": "S101", "type": "Stütze", "section_type": "rec", "dimension": "10/20", "units": "cm"},
        {"position": "T101", "type": "Träger", "section_type": "rec", "dimension": "120/240"},
        {"position": "S102", "type": "Stütze", "section_type": "", "dimension": ""},
    ]
    assert store.last_section_of_type(rows, "Stütze") == {"section_type": "rec", "dimension": "10/20", "units": "cm"}, "the last that states one"
    assert store.last_section_of_type(rows, "Träger") == {"section_type": "rec", "dimension": "120/240", "units": ""}
    assert store.last_section_of_type(rows, "Decke/Dach") is None
    copy = store.section_of(rows[0])
    copy["dimension"] = "12/24"
    assert rows[0]["dimension"] == "10/20", "a copy, never a link"
    assert store.section_of({"position": "X1"}) is None


def test_the_types_and_symbols_are_the_ste_legend():
    from domains import plan
    names = store.types()
    assert names[0] == plan.TYPES["V"]["names"][None] and "Träger" in names and "Decke/Dach" in names and "Betondecke" in names
    assert len(names) == len(set(names))
    assert store.symbol_for("Träger") == "T" and store.symbol_for("Betondecke") == "B" and store.symbol_for("") == ""
    assert store.symbol_for("Pfette", [{"position": "P001", "type": "Pfette"}]) == "P", "a type of the engineer's own: the letter of a position of it"
    assert store.type_for("S") == "Stütze" and store.type_for("D") == "Scheibenbildung Decke/Dach", "a letter with subtypes: its first"
    assert store.type_for("P", [{"position": "P001", "type": "Pfette"}]) == "Pfette" and store.type_for("Q") == ""


def test_a_blank_position_is_made_for_the_plan_tab():
    here = _here("store_blank")
    row = store.create_blank(here, None, "Träger")
    assert row == store.normalise({"position": "T002", "type": "Träger"}), row
    assert store.create_blank(here, "T002", "Träger") == row and len(store.read(here)) == 4, "asking twice writes once"
    row = store.create_blank(here, None, "Stütze", floor="1")
    assert row["position"] == "S101"
    assert json.loads(store.path(here).read_text(encoding="utf-8"))["positions"][-1] == {"position": "S101", "type": "Stütze"}


# --- the workbook of a project made before v0.114 ------------------------------------------

def test_a_workbooks_rows_are_mapped_and_its_gaps_said_not_refused():
    rows, notes = store.from_sheet([
        {"position": "T001", "type": "Träger", "resistance class": "C24", "section type": "rec", "dimension": "100/200", "units": "mm", "moisture class": 2.0},
        {"position": "S001", "type": None, "Wood moisture class": "1", "dimension": 180},
        {"position": "t001", "type": "Träger"},
        {"position": None, "type": "Träger", "material": "Fichte"},
        {"position": "X9", "type": None},
        {"position": "D001", "type": "Decke/Dach", "moisture class": 7},
        {"position": None, "type": None},
    ])
    assert [row["position"] for row in rows] == ["T001", "S001", "D001"]
    assert rows[0]["strength_class"] == "C24" and rows[0]["section_type"] == "rec" and rows[0]["moisture_class"] == 2
    assert rows[1]["type"] == "Stütze" and rows[1]["moisture_class"] == 1 and rows[1]["dimension"] == "180", "the older header still read"
    assert rows[2]["moisture_class"] is None
    codes = [note.code for note in notes]
    assert codes == ["positions_import_type", "positions_import_twice", "positions_import_no_id", "positions_import_no_type", "positions_import_moisture"], codes
    assert all(isinstance(note, engine.Note) for note in notes)


def test_the_migration_moves_the_workbook_into_the_archives_and_is_idempotent():
    here = _here("store_migrate")
    store.path(here).unlink()
    book = _workbook(store.workbook(here), [{"position": "T001", "type": "Träger", "material": "Fichte", "moisture class": 1},
                                           {"position": "S001", "type": "Stütze"}])
    note = store.migrate(here)
    assert note is not None and note.code == "positions_migrated" and note.args["count"] == 2, note
    assert not book.exists(), "the workbook is not read any more, so it does not stay where it was read from"
    archived = list(here.archives_dir.glob("*-Positionen.xlsx"))
    assert len(archived) == 1 and archived[0].stat().st_size > 0, "moved, never deleted"
    assert [row["position"] for row in store.read(here)] == ["T001", "S001"]
    assert store.migrate(here) is None, "a second migration finds positions.json and does nothing"
    # a workbook put back beside a positions.json is not read again
    shutil.copy2(archived[0], book)
    assert store.migrate(here) is None and book.exists()
    # and a workbook that cannot be read writes nothing, for the next opening to try again
    store.path(here).unlink()
    book.write_bytes(b"not a workbook")
    _refused("unreadable", store.migrate, here)
    assert not store.path(here).exists() and book.exists()


def test_import_merges_adding_what_is_missing_or_replacing():
    rows = [store.normalise({"position": "T001", "type": "Träger", "material": "Fichte"})]
    imported = [store.normalise({"position": "t001", "type": "Träger", "material": "Lärche"}), store.normalise({"position": "S001", "type": "Stütze"})]
    merged, added, replaced = store.merge(rows, imported)
    assert added == ["S001"] and replaced == [] and merged[0]["material"] == "Fichte"
    merged, added, replaced = store.merge(rows, imported, replace=True)
    assert added == ["S001"] and replaced == ["t001"] and merged[0]["material"] == "Lärche" and len(merged) == 2


# --- the export and the report -------------------------------------------------------------

def test_a_section_that_does_not_read_is_a_warning_and_a_blank_position_exports():
    import export_all
    positions, notes = export_all._positions([
        {"position": "Z001", "type": "Träger"},
        {"position": "Z002", "type": "Stütze", "section_type": "ROR", "dimension": "101.6"},
        {"position": "Z003", "type": "Stütze", "section_type": "Holz", "dimension": "3"},
        {"position": "Z004", "type": "Stütze", "section_type": "rec"},
        {"position": "Z005", "type": "Stütze", "moisture_class": 7},
    ])
    assert [p["section"] for p in positions] == [None] * 5 and positions[4]["moisture_class"] is None
    assert [n.code for n in notes] == ["position_section"] * 3 + ["position_moisture"], [n.code for n in notes]
    assert positions[0]["schema_version"] == export_all.POSITION_SCHEMA_VERSION == "1.4"
    for record in positions:
        validation.validate_structure(record, "position")
    validation.validate_structure({"schema_version": "1.4", "position": "Z009", "type": "Träger", "domains": {}}, "position")


def test_the_export_writes_a_project_with_no_workbook_and_check_migrates_nothing():
    """`woodcore-cli check` writes nothing: it reads a workbook a migration would read, and moves nothing."""
    import export_all
    here = _here("store_check")
    store.path(here).unlink()
    book = _workbook(store.workbook(here), [{"position": "T001", "type": "Träger"}])
    with project.using(here):
        try:
            export_all.run(dry=True)
        except validation.ValidationError as exc:
            # what the copy lacks further on -- a plan's background, which conftest.project_copy leaves out -- is not this test's; the positions are read before it
            assert "Positionen" not in str(exc) and store.FILE not in str(exc), exc
    assert book.exists() and not store.path(here).exists(), "a dry run moved the workbook"


def _planted(name, records):
    """A Typst root of the test's own, the example's engine and export, `records` planted as positions."""
    root = conftest.tmpdir(name)
    for f in ("commands.typ", "config.typ"):   # where the project keeps them: 000-Typst-Config/ laid out by folders (issue #64)
        source = conftest.PROJECT.config_dir / f
        target = root / source.relative_to(conftest.PROJECT.typst_root)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
    generated = conftest.PROJECT.generated_dir
    relative = generated.relative_to(conftest.TYPST_ROOT)
    shutil.copytree(generated, root / relative, copy_function=shutil.copyfile, ignore=shutil.ignore_patterns("out", "docs"))
    data = root / relative / "data"
    ids = json.loads((data / "positions-index.json").read_text(encoding="utf-8"))["positions"]
    for record in records:
        (data / "positions" / f"{record['position']}.json").write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
        ids.append(record["position"])
    (data / "positions-index.json").write_text(json.dumps({"schema_version": "1.0", "positions": ids}), encoding="utf-8")
    return root


def _compiles(root, body):
    if not conftest.have_typst():
        raise Skip("typst is not on the path")
    probe = root / "probe.typ"
    probe.write_text('#import "/000-Typst-Config/commands.typ": *\n' + body + "\n", encoding="utf-8")
    out = subprocess.run(["typst", "compile", "--root", str(root), str(probe), str(root / "probe.pdf")], capture_output=True, text=True, encoding="utf-8")
    return out.returncode == 0, conftest.panic_message(out.stderr) or out.stderr


def test_a_blank_position_goes_through_the_report_without_a_panic():
    """A position the plan tab made (id and type only) is printed and written down as not calculable -- and a record carrying no more than the schema requires is read with defaults, every optional key read with `.at(…, default: …)`."""
    import export_all
    (blank,), _notes = export_all._positions([{"position": "Z001", "type": "Träger"}])
    bare = {"schema_version": "1.4", "position": "Z002", "type": "Stütze", "domains": {}}
    root = _planted("store_typst", [blank, bare])
    every = '("type", "material", "strength-class", "section-type", "dimension", "units", "moisture-class", "length", "descr")'
    for body in ('#position("Z001", category-show: %s)' % every, '#position("Z002", category-show: %s)' % every, '#for id in ids("position") { position(id) }'):
        ok, message = _compiles(root, body)
        assert ok, (body, message)


def test_a_blank_position_warns_a_schema_and_a_connector_and_stops_neither():
    from domains import connector, schema
    import export_all
    here = conftest.PROJECT
    records = {}
    for pid in ("T001", "S001", "S002"):
        (record,), _ = export_all._positions([{"position": pid, "type": "Träger"}])
        records[pid] = record
    details = schema.read_sources(here.details_dir)
    _schemas, _files, schema_warnings = schema.prepare(here.details_dir, records)
    _data, _figures, warnings = connector.prepare(here.connectors_dir, records, details)
    assert schema_warnings or warnings, "a position with no section and no strength class must be said somewhere"


def test_a_project_laid_out_by_folders_keeps_its_positions_in_positions_and_a_workbook_beside_them():
    """Issue #64 (D-24): positions.json, the source, in `Positions/`, and `Positions.xlsx` written beside it at every change -- the sheet and the headers of Positionen.xlsx, so a copy of it imports again -- and never read: a workbook deleted or edited changes nothing of what is read."""
    import documents
    import openpyxl
    here = project.write(conftest.tmpdir("positions_folders") / "P.woodcore", "P", documents.layout_for("fr"))
    assert store.path(here) == here.root / "Positions" / "positions.json"
    rows = [{"position": "S101", "type": "Pfette", "material": "Fichte", "strength_class": "C24", "section_type": "rect", "dimension": "120/240", "units": "mm", "moisture_class": 1, "description": "Pfette Achse 1"}]
    store.write(here, rows)
    book = store.workbook_path(here)
    assert book == here.root / "Positions" / "Positions.xlsx" and book.is_file()
    sheet = openpyxl.load_workbook(book)[store.SHEET]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0] == [store.HEADERS[key] for key in store.KEYS]
    assert values[1][:3] == ["S101", "Pfette", "Fichte"] and values[1][7] == 1
    book.write_bytes(b"not a workbook")
    assert store.read(here)[0]["position"] == "S101", "the workbook is never read"
    # a project laid out before writes none
    alone = project.write(conftest.tmpdir("positions_alone") / "A.woodcore", "A")
    store.write(alone, rows)
    assert store.workbook_path(alone) is None and not list(alone.root.rglob("Positions.xlsx"))
