"""Builds a minimal Belastungsannahmen workbook, for a test that needs one.

A test must not read the engineer's own workbook. Two reasons, and the second is the one that matters: a test pointing at `02_Belastungsannahmen.xlsm` fails the day a row is edited, and a test that fails for a reason nobody caused gets deleted rather than fixed. The other is that a defect has to be *put* into a sheet to check that the export refuses it, and nothing here may write to the engineer's files.

So the sheets are built here, with only what the reader actually reads -- the layer table of `domains/dead_load.py` and the material library it looks up against. Everything else the real workbook carries (the wind block, the Dropdowns, the Objekt notes) is left out: a fixture that reproduced the whole workbook would be a second workbook to maintain.

⚠️ openpyxl writes no formula *cache*. A cell this pipeline reads by value must therefore be written as a value; only column I, which is read as a formula (`_check_weight_formula`), may carry one.
"""

from openpyxl import Workbook

MATERIAL_SHEET = "Material (hier ergänzen)"

# One row of the layer table, as `build()` takes it.
#   kind     "gk" (Eigenlast) or "ak" (Auflast)
#   ebene    the layer's role, which keys the library with the material
#   material the name the library is looked up by
#   d, b, e  Dicke, Breite, Dist. in m
#   weight   what column I actually holds: None writes the real lookup formula,
#            a number writes that number over it -- which is the defect
#            `_check_weight_formula` exists to refuse.
#   load     a load typed straight into column K of a row that names no material, as the Belastungsannahmen template 3.2 does for
#            «Installationen und Innenwände» -- counted by the sheet, and until the audit of 260916 not by the export.
LAYER = ("kind", "ebene", "material", "d", "b", "e", "weight", "load")


def layer(kind, ebene, material, d=1, b=1, e=1, weight=None, load=None):
    return dict(zip(LAYER, (kind, ebene, material, d, b, e, weight, load)))


def direct(kind, ebene, load):
    """A row with no material and its load typed in column K."""
    return layer(kind, ebene, None, d=None, b=None, e=None, load=load)


def build(path, sheet="EL+AL-Test", layers=(), library=(), chosen=None):
    """Writes a workbook at `path` and returns it.

    `library` is [(ebene, material, mass, unit), ...]; `chosen` the *Gewählt* permanent load, which defaults to the computed sum rounded up to the next 0.1 -- the normal case, which must pass.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    lib = {}
    for ebene, material, mass, unit in library:
        lib[(ebene, material)] = (mass, unit)

    # The layer table: the header row `dead_load` finds by its label, the layers, then the Total line that ends it.
    header = 6
    ws[f"C{header}"] = "Ebene"
    row = header + 1
    total_load = 0.0
    for spec in layers:
        if spec.get("load") is not None and not spec["material"]:
            ws[f"B{row}"] = spec["kind"]
            ws[f"D{row}"] = spec["ebene"]
            ws[f"K{row}"] = spec["load"]
            total_load += spec["load"]
            row += 1
            continue
        mass, unit = lib.get((spec["ebene"], spec["material"]), (0, "kg/m3"))
        ws[f"B{row}"] = spec["kind"]
        ws[f"D{row}"] = spec["ebene"]
        ws[f"E{row}"] = spec["material"]
        ws[f"F{row}"] = spec["d"]
        ws[f"G{row}"] = spec["b"]
        ws[f"H{row}"] = spec["e"]
        # column I is the sheet's own XLOOKUP. A real one is a formula; `weight` writes a number over it.
        if spec["weight"] is None:
            ws[f"I{row}"] = f'=IF(E{row}="","",_xlfn.XLOOKUP($E{row},Tabelle2[Material],Tabelle2[Masse]))'
        else:
            ws[f"I{row}"] = spec["weight"]
        # a fixture may state a Dist. of 0 on purpose, to check that the reader refuses it
        if spec["e"]:
            total_load += mass * spec["d"] * spec["b"] / (spec["e"] * 100.0)
        row += 1

    total_row = row + 1
    ws[f"J{total_row}"] = "Total"

    chosen_row = total_row + 2
    ws[f"D{chosen_row}"] = "Ständige Lasten"
    if chosen is None:
        # what an engineer does: round the computed sum up to the next tenth
        import math
        chosen = math.ceil(total_load * 10) / 10
    ws[f"J{chosen_row}"] = chosen

    lib_ws = wb.create_sheet(MATERIAL_SHEET)
    lib_ws["A3"], lib_ws["B3"], lib_ws["C3"], lib_ws["D3"] = "Ebene", "Material", "Masse", "Einheit"
    for i, (ebene, material, mass, unit) in enumerate(library, start=4):
        lib_ws[f"A{i}"], lib_ws[f"B{i}"] = ebene, material
        lib_ws[f"C{i}"], lib_ws[f"D{i}"] = mass, unit

    wb.save(path)
    return {"path": path, "sheet": sheet, "computed": total_load, "chosen": chosen}


# --- the snow and wind sheets (spec 2026-10-21-restes-issue-69, D-09) -----------------------------------------
# The workbook's own Dropdowns, as the engineer's template carries them: the labels the SL-*/WL-* sheets pick from, and the numbers each stands for (held to SIA 261 by the readers).
DROPDOWNS = {
    "B24:C26": [("normale Windexposition", 1), ("dem Wind stark ausgesetzte Lage", 0.8), ("vor Wind geschützte Lage", 1.2)],
    "B30:C34": [(-200, -200), ("± 0", 0), ("+200", 200), ("+400", 400), ("+500", 500)],
    "I44:I48": [("Flachdach",), ("GiebeldachLM1",), ("GiebeldachLM2",), ("SheddachLM1",), ("SheddachLM2",)],
    "B91:D94": [("II Seeufer", 300, 0.16), ("IIa grosse Ebene", 380, 0.19), ("III Ortschaften, freies Feld", 450, 0.23), ("IV grossflächige Stadtgebiete", 526, 0.3)],
}


def _cells(rng):
    from openpyxl.utils import range_boundaries
    c0, r0, c1, r1 = range_boundaries(rng)
    return c0, r0, c1, r1


def build_loads(path, snow=(), wind=(), dropdowns=None):
    """Writes a load workbook at `path` with a `Dropdowns` sheet and the given SL-*/WL-* sheets, each {"sheet", "inputs": {cell: value}, "excel": {cell: value}} -- the cells the engineer enters, and those Excel computed (written as values: openpyxl writes no formula cache). tests/fixtures/load-sheets-reference.json holds such sheets, read from the engineer's workbooks."""
    wb = Workbook()
    drop = wb.active
    drop.title = "Dropdowns"
    for rng, rows in (dropdowns or DROPDOWNS).items():
        c0, r0, _c1, _r1 = _cells(rng)
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                drop.cell(row=r0 + i, column=c0 + j, value=value)
    for item in list(snow) + list(wind):
        ws = wb.create_sheet(item["sheet"])
        for cell, value in {**item.get("excel", {}), **item.get("inputs", {})}.items():
            ws[cell] = value
    wb.save(path)
    return path
