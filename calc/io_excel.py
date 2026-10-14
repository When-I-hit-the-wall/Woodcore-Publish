"""Thin openpyxl reader: the only module that touches an Excel workbook.

Two shapes of input, two readers (docs/architecture/principles.md §2, docs/domain/loads.md §3.3bis):

* `read_workbook()` -- Positionen.xlsx, whose sheets are tables: each one becomes a list of row dicts keyed by header.
* `open_cells()` -- a Belastungsannahmen workbook, whose sheets are the forms the engineer fills in. Values sit at fixed cell addresses on a sheet whose *name* identifies the case ("SL-Dach1", "EL+AL-Fassade"), so it is read cell by cell, sheet by sheet.

The workbook's own note (sheet `Objekt`) says its sheets are meant to be copied -- "Blätter können kopiert werden", "Bauteilname wird von Blattname übernommen". `sheet_names()` is what turns that into data: a domain declares a prefix, every sheet carrying it becomes one instance of that calculation, and copying a sheet in Excel adds a case to the report without touching any code. `find_workbooks()` does the same one level up: which files are load workbooks is decided by their name, so a project's `02_Belastungsannahmen.xlsm` is read without being named anywhere.
"""

import warnings

import validation


def _load(path, **kwargs):
    """openpyxl warns loudly about the .xlsm's data-validation extension, its drawings and its wmf images -- none of which we read. Silenced so a clean export run prints nothing but its own progress lines.
    """
    # openpyxl (and numpy behind it) imported here, when a workbook is read, never at import: an export of a project with no workbook does without its two seconds (docs/architecture/performance.md, v0.117)
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return load_workbook(path, **kwargs)


def read_workbook(path):
    wb = _load(path, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sheets[name] = []
            continue
        headers = [str(h).strip() for h in rows[0]]
        sheet_rows = []
        for row in rows[1:]:
            if all(cell is None for cell in row):
                continue
            sheet_rows.append(dict(zip(headers, row)))
        sheets[name] = sheet_rows
    return sheets


def _clean(value):
    """Excel keeps the trailing spaces an operator typed into a cell (and a dropdown's own label may carry one); every value we read is compared against a fixed set or printed in the report, so strip them here once.
    """
    return value.strip() if isinstance(value, str) else value


class CellReader:
    """Random-access reader for one workbook's cached cell values.

    openpyxl never evaluates formulas: it returns the value Excel cached the last time the file was saved. A workbook that has never been opened in Excel therefore hands back None for every computed cell -- reported here as a blocking error naming the cell, rather than silently exported as an empty load.
    """

    def __init__(self, path):
        self.path = path
        self._wb = _load(path, data_only=True)
        # Parsed on first use only: the formulas are a second full read of the workbook, and only the shape-coefficient helper (domains/snow.py) needs them.
        self._formulas = None

    def address(self, sheet, ref=None):
        """Where a value was read, as provenance and error messages both print it: `02_Belastungsannahmen.xlsm!SL!H17`, or the sheet alone without a `ref`. Several load workbooks may be read in one export, so a sheet name no longer says which file it is in -- every address names its workbook, and is built here so none can forget to.
        """
        where = f"{self.path.name}!{sheet}"
        return where if ref is None else f"{where}!{ref}"

    def formula(self, sheet, ref):
        """What a cell was *told* to compute, as authored -- `=IF(F21="Flachdach",Dropdowns!F61,Dropdowns!F64)`.

        `value()` reads what Excel computed; this reads the formula behind it, which is the only thing that says *which* sheet's numbers went into that value. A helper cell shared by every sheet of a domain cannot be told apart from its value alone, so the sheet it actually reads is checked against the sheet being exported (docs/domain/loads.md §3.3bis). Returns the literal for a cell holding one, None for an empty cell.
        """
        if self._formulas is None:
            self._formulas = _load(self.path, data_only=False)
        if sheet not in self._formulas.sheetnames:
            raise validation.ValidationError(
                f"{self.path.name}: no sheet '{sheet}' -- found: "
                f"{', '.join(self._formulas.sheetnames)}"
            )
        raw = self._formulas[sheet][ref].value
        # An array formula comes back as an object carrying its own text.
        return getattr(raw, "text", raw)

    def sheet_names(self, prefix):
        """Every sheet whose name starts with `prefix`, in workbook order -- the sheets of one calculation domain (docs/domain/calculation-domains.md §3.3). Returns an empty list rather than failing: a project may simply have no facade, or no snow.
        """
        return [name for name in self._wb.sheetnames if name.startswith(prefix)]

    def max_row(self, sheet):
        """Last row openpyxl sees content on -- the bound a lookup table of unknown length is scanned to (the material library).
        """
        return self._sheet(sheet).max_row

    def _sheet(self, sheet):
        if sheet not in self._wb.sheetnames:
            raise validation.ValidationError(
                f"{self.path.name}: no sheet '{sheet}' -- found: "
                f"{', '.join(self._wb.sheetnames)}"
            )
        return self._wb[sheet]

    def value(self, sheet, ref, required=True):
        """One cell, as Excel last computed it. `required=False` returns None for a cell the sheet leaves empty (a load model's second shape coefficient, say) instead of failing.
        """
        raw = _clean(self._sheet(sheet)[ref].value)
        if raw == "":
            raw = None
        if raw is None and required:
            raise validation.ValidationError(
                f"{self.address(sheet, ref)} is empty -- open the workbook in "
                f"Excel, check that cell, and save, so its computed value is cached"
            )
        if isinstance(raw, str) and raw.startswith("#"):
            raise validation.ValidationError(
                f"{self.address(sheet, ref)} holds the Excel error {raw} -- "
                f"fix it in the workbook and save before exporting"
            )
        return raw

    def at(self, sheet, column, row, required=True):
        """`value()` addressed as (column letter, row number) -- the shape a row-scanning reader needs, since the row is found rather than known in advance.
        """
        return self.value(sheet, f"{column}{row}", required=required)

    def find_row(self, sheet, column, text, first=1, last=200):
        """Row number of the first cell in `column` holding `text`.

        The layer tables of the EL+AL sheets do not end at a fixed row -- a facade has fewer layers than a roof, so its `Total` line sits two rows higher (docs/domain/loads.md §3.3ter). Anything below such a table is therefore located by its own label instead of by an address that would silently point at the wrong row.
        """
        ws = self._sheet(sheet)
        for row in range(first, last + 1):
            if _clean(ws[f"{column}{row}"].value) == text:
                return row
        raise validation.ValidationError(
            f"{self.address(sheet)}: no cell in column {column} holds "
            f"'{text}' (rows {first}-{last}) -- the sheet's layout changed, "
            f"follow it in the domain module that reads it"
        )

    def pairs(self, sheet, ref):
        """A two-column range as [(left, right), ...] -- the shape of this workbook's `Dropdowns` lookup tables (label, value). Rows that are blank on either side are skipped.
        """
        out = []
        for row in self._sheet(sheet)[ref]:
            left, right = (_clean(c.value) for c in row[:2])
            if left is None or right is None:
                continue
            out.append((left, right))
        return out

    def column(self, sheet, ref):
        """A one-column range as a list, blanks skipped."""
        return [
            _clean(row[0].value)
            for row in self._sheet(sheet)[ref]
            if _clean(row[0].value) not in (None, "")
        ]

    def rows(self, sheet, columns, first, last):
        """Rows `first`..`last` as [(row_number, {name: value}), ...], `columns` mapping a name to a column letter. Values come back raw (None for an empty cell) -- a row-scanning reader decides for itself which rows count, so nothing is required here.
        """
        ws = self._sheet(sheet)
        return [
            (row, {name: _clean(ws[f"{col}{row}"].value) for name, col in columns.items()})
            for row in range(first, last + 1)
        ]


def open_cells(path):
    if not path.exists():
        raise validation.ValidationError(f"missing input workbook: {path}")
    return CellReader(path)


# Excel keeps a lock file beside a workbook for as long as it is open -- `~$Belastungsannahmen.xlsm`, a few bytes of owner name, not a workbook. It matches any name pattern its workbook does, and the engineer usually has that workbook open while building.
_LOCK_FILE_PREFIX = "~$"


def find_workbooks(folder, contains, suffixes):
    """Every workbook directly in `folder` whose name contains `contains` and ends in one of `suffixes`, ordered by file name.

    Both tests ignore case, as Windows does -- `belastungsannahmen.XLSX` is the same file to the engineer. Subfolders are not searched: a copy filed away in an archive folder is exactly what must not turn into a second set of loads. The order is the file name's, so an office's numbering (`01_...`, `02_...`) is the order the workbooks' cases come out in.
    """
    needle = contains.casefold()
    wanted = {s.casefold() for s in suffixes}
    return sorted(
        (
            p for p in folder.iterdir()
            if p.is_file()
            and needle in p.name.casefold()
            and p.suffix.casefold() in wanted
            and not p.name.startswith(_LOCK_FILE_PREFIX)
        ),
        key=lambda p: p.name.casefold(),
    )
