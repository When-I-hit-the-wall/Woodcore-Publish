"""Generates docs/typst-reference.xlsx: the browsable catalog of every command this project defines (parsed from the /// doc comments of lib/, docs/typst/api-conventions.md §10bis) and every valid show-* argument combination for the current project (parsed from data/*.json, docs/typst/reference-workbook.md §10ter).

Since v0.64 it reads **all three files commands.typ re-exports** -- calc-engine.typ, formula.typ and i18n.typ -- and every top-level `#let` in them, not only the ones whose name is followed by a parenthesis. Before that it read calc-engine.typ alone and could not see a binding like `#let report-level = state(...)`, so var, eq, verify, val, unit, var-domain, var-unit, the ten commands of i18n.typ and the four settings of the document were in no catalog at all. It also carries the categories Woodcore's commands tab lists them under, and the worked examples of each (section 10sexies) -- so the workbook, the window and Tinymist's hover all read the same text.

Every show-* call is addressed the same way -- `(domain, instance)`, then whatever it looks up inside -- so every sheet below leads with those two columns. An instance is a workbook sheet for a calculation domain (data/calculations.json), a position id for a domain attached to a position (data/positions/<id>.json), or a connection's id under the domain "connection" (data/connections.json); one `iter_sources()` enumerates all three, so no sheet can list one kind and forget another. The Connections sheet adds the call no other kind has, show-polar-moment, and names the schema each is drawn on; the Connectors sheet lists the connectors (data/connectors.json), the polar-moment calculations drawn on the same schema as each, and the show-connector call that prints it.

Like everything under data/, this workbook is generated and must never be hand-edited -- calc/export_all.py regenerates it on every run. Can also be run standalone right after editing a /// block, without rerunning the whole data pipeline.
"""

import io
import json
import functools
import re
import sys
import zipfile
from pathlib import Path
try:
    from build_reference_lib import (  # split out of this module (issue #30): every name it had is still importable from it
        BULLET_NAME_RE,
        BULLET_RE,
        CATEGORIES,
        DOC_LINE_RE,
        FENCE_RE,
        INTERNAL,
        LET_FN_RE,
        LET_RE,
        LIB_FILES,
        MULTI_BULLET_RE,
        RUNNABLE_FENCE,
        VARIADIC,
        category_of,
        lib_bindings,
        lib_records,
        parse_bindings,
        parse_doc_block,
        parse_doc_comments,
        parse_param_docs,
        read_binding,
        read_signature,
        split_params,
        value_signature,
    )
except ImportError:   # the export imports build_reference as tools.build_reference, calc/tools not on its path
    from tools.build_reference_lib import (  # split out of this module (issue #30): every name it had is still importable from it
        BULLET_NAME_RE,
        BULLET_RE,
        CATEGORIES,
        DOC_LINE_RE,
        FENCE_RE,
        INTERNAL,
        LET_FN_RE,
        LET_RE,
        LIB_FILES,
        MULTI_BULLET_RE,
        RUNNABLE_FENCE,
        VARIADIC,
        category_of,
        lib_bindings,
        lib_records,
        parse_bindings,
        parse_doc_block,
        parse_doc_comments,
        parse_param_docs,
        read_binding,
        read_signature,
        split_params,
        value_signature,
    )
try:
    from build_reference_calls import (  # split out of this module (issue #30): every name it had is still importable from it
        ROOT,
        Raw,
        TYPST_GLOBALS,
        _NONE,
        _args,
        _caption,
        _given,
        _same,
        _signatures,
        call_calc_block_plan,
        call_check_member_loop,
        call_check_steel_loop,
        call_command,
        call_connector_loop,
        call_expression,
        call_extract_loop,
        call_find,
        call_fire_plan_loop,
        call_instance_loop,
        call_journal_entries,
        call_plan_loop,
        call_schema_loop,
        call_show_annexes,
        call_show_calc,
        call_show_check,
        call_show_check_member,
        call_show_check_steel,
        call_show_connector,
        call_show_connector_resistance,
        call_show_extract,
        call_show_fire_plan,
        call_show_input,
        call_show_journal,
        call_show_line_load_summary,
        call_show_live_load,
        call_show_load_summary,
        call_show_plan,
        call_show_plan_legend,
        call_show_polar_moment,
        call_show_position_list,
        call_show_schema,
        call_show_static_model,
        call_static_model_loop,
        call_val,
        call_value_of,
        call_var,
        call_var_of,
        call_var_registry,
        registry_key,
        signature_defaults,
        typst_ident,
        typst_number,
        typst_str,
        typst_value,
    )
except ImportError:   # the export imports build_reference as tools.build_reference, calc/tools not on its path
    from tools.build_reference_calls import (  # split out of this module (issue #30): every name it had is still importable from it
        ROOT,
        Raw,
        TYPST_GLOBALS,
        _NONE,
        _args,
        _caption,
        _given,
        _same,
        _signatures,
        call_calc_block_plan,
        call_check_member_loop,
        call_check_steel_loop,
        call_command,
        call_connector_loop,
        call_expression,
        call_extract_loop,
        call_find,
        call_fire_plan_loop,
        call_instance_loop,
        call_journal_entries,
        call_plan_loop,
        call_schema_loop,
        call_show_annexes,
        call_show_calc,
        call_show_check,
        call_show_check_member,
        call_show_check_steel,
        call_show_connector,
        call_show_connector_resistance,
        call_show_extract,
        call_show_fire_plan,
        call_show_input,
        call_show_journal,
        call_show_line_load_summary,
        call_show_live_load,
        call_show_load_summary,
        call_show_plan,
        call_show_plan_legend,
        call_show_polar_moment,
        call_show_position_list,
        call_show_schema,
        call_show_static_model,
        call_static_model_loop,
        call_val,
        call_value_of,
        call_var,
        call_var_of,
        call_var_registry,
        registry_key,
        signature_defaults,
        typst_ident,
        typst_number,
        typst_str,
        typst_value,
    )


CATEGORY_KEYS = tuple(key for key, _names in CATEGORIES) + (INTERNAL,)

PUBLIC_NAMES = tuple(name for _key, names in CATEGORIES for name in names)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def load_positions(data_dir):
    return [load_json(p) for p in sorted((data_dir / "positions").glob("*.json"))]


# --- The Typst calls a reader pastes (docs/typst/reference-workbook.md §10ter, docs/ui/window.md §10quater) ---
# The one place a show-* call is spelled out as text. This workbook's example_call columns and Woodcore's window (tools/woodcore_gui.py) both go through these functions, so the same call cannot come out two different ways -- and when a signature in lib/calc-engine.typ changes, this is the one place to follow it.

# Where every command is imported from, as seen from Sections/*.typ, which is where these calls get pasted. commands.typ re-exports the whole of lib/calc-engine.typ (docs/typst/api-conventions.md §10bis), so its one import line serves every call below, and any command added to the engine later, with no list of names to keep up to date.
COMMANDS_IMPORT_PATH = "../commands.typ"

# The domain a connection's record is read under (docs/domain/connections.md §3.3quinquies): `show-calc("connection", "VU-MP-1")`.
CONNECTION_DOMAIN = "connection"

# The domain read from a norm table rather than from a workbook, and the engine that says so (docs/domain/loads.md §3.3septies): the Nutzlasten of SIA 261's Tabelle 8, whose instances are its categories and which has two calls of its own.
LIVE_LOAD_DOMAIN = "live-load"
NORM_ENGINE = "norm"


def load_statics(data_dir):
    """data/statics.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "statics.json"
    return load_json(path) if path.exists() else {"models": {}}


def build_statics_rows(statics):
    """One row per static model: what it is built of, what its loads show and where their numbers came from, and the call that prints it."""
    rows = []
    for sid, record in statics.get("models", {}).items():
        sources = sorted({load["from"] for load in record["loads"] if load.get("from")})
        rows.append((
            sid,
            record.get("label", ""),
            len(record["points"]),
            len(record["members"]),
            len(record["supports"]),
            len(record["loads"]),
            len(record.get("dimensions", ())),
            "; ".join(load["text"] for load in record["loads"] if load["text"]),
            "; ".join(f"{dim['id']} = {dim['text']}" for dim in record.get("dimensions", ())),
            ", ".join(sources),
            ", ".join(record["unresolved"]),
            call_show_static_model(sid),
        ))
    return rows


def load_schemas(data_dir):
    """data/schemas.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "schemas.json"
    return load_json(path) if path.exists() else {"schemas": {}}


def build_schemas_rows(schemas):
    """One row per detail schema: what it is built of, and the call that prints it."""
    rows = []
    for sid, record in schemas.get("schemas", {}).items():
        rows.append((
            sid, record.get("label", ""),
            len(record["elements"]), len(record["connectors"]), len(record["aids"]),
            len(record["dimensions"]),
            ", ".join(e["id"] for e in record["elements"]),
            ", ".join(c["id"] for c in record["connectors"]),
            ", ".join(record["unresolved"]),
            call_show_schema(sid),
        ))
    return rows


def load_plans(data_dir):
    """data/plans.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "plans.json"
    return load_json(path) if path.exists() else {"plans": {}}


def load_fire_plans(data_dir):
    """data/fire.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "fire.json"
    return load_json(path) if path.exists() else {"plans": {}}


def build_fire_plans_rows(fire):
    """One row per fire-protection plan: the position plan it follows or the document page it is drawn on, what it draws, and the call that prints it."""
    rows = []
    for pid, record in fire.get("plans", {}).items():
        background = record.get("background") or {}
        origin = record.get("plan") or (background.get("file") or "")
        scale = record.get("scale", {})
        rows.append((
            pid, record.get("label", ""), "follows " + record["plan"] if record.get("plan") else "own frame", origin,
            f"1:{scale['ratio']:g}" if scale.get("ratio") else ("kalibriert" if scale else ""),
            len(record.get("zones", ())), len(record.get("walls", ())), len(record.get("doors", ())),
            len(record.get("routes", ())) + len(record.get("exits", ())) + len(record.get("widths", ())),
            call_show_fire_plan(pid),
        ))
    return rows


def build_plans_rows(plans):
    """One row per position plan: what it is drawn on, what it shows, the labels Positionen.xlsx did not have at the export, and the call that prints it."""
    rows = []
    for pid, record in plans.get("plans", {}).items():
        background = record.get("background")
        origin = f"{background['origin']} (S. {background['origin_page']})" if background else ""
        scale = record.get("scale", {})
        rows.append((
            pid, record.get("label", ""), record.get("floor", ""), origin,
            f"1:{scale['ratio']:g}" if scale.get("ratio") else ("kalibriert" if scale else ""),
            len(record["elements"]), len(record["axes"]), ", ".join(record["labels"]), ", ".join(record["unknown"]),
            call_show_plan(pid),
        ))
    return rows


# --- The two calls the command tab copies (docs/typst/command-docs.md §10sexies) ---
# «Complet» writes every argument out with its default, so a reader deletes what he does not want instead of looking the rest up; «minimal» is the same call reduced to what must be given. Both are built from the signature parsed out of lib/, so an argument added to a command reaches both with no edit here -- the argument of the call_* functions above, applied to calls nobody wrote by hand.

TRAILING_BODY = "body"  # a content block, written after the parentheses: #calc-block("T001")[ ... ]

# A binding that is not a function has no call to derive: what one does with a state is written out here.
WRITTEN_CALLS = {
    "report-level": ("#report-level.update(2)",) * 2,
    "uniform-level": ("#uniform-level.update(2)",) * 2,
    "number-digits": ("#number-digits.update(3)",) * 2,
    "report-language": ("#context report-language.get()",) * 2,
    # a show rule is applied, never called: its one argument is the document it is handed (v0.82)
    "woodcore-document": ("#show: woodcore-document",) * 2,
    "caption-rules": ("#show: caption-rules",) * 2,
    # its one argument is the document it cites, so a call that leaves it out is a call that refuses -- there is no default a citation could have (v0.99)
    "cite-norm": ('#cite-norm(norm: "SIA 265", norm-ref: "4.2.2.1")',) * 2,
    # a show rule, applied to the sheet it is the top of (v0.114)
    "sheet": ('#show: sheet.with(format: "A3", landscape: true, margin: 10mm)', "#show: sheet"),
    # its one argument is a variable a report or its export carries, and SIA 261 Tabelle 8's live loads are in every export, a new project's included
    "doc-value": ('#doc-value("q_k", domain: ("live-load", "B"))',) * 2,
    # these read the report's language, so only inside `context` (periodic audit 260916)
    "report-lang": ("#context report-lang()",) * 2,
    "t-str": ('#context t-str("governing")',) * 2,
    "tl-str": ('#context tl-str("Bezugshöhe")',) * 2,
    "t-label": ('#context t-label("governing")',) * 2,
    "t-pattern": ('#context t-pattern("connection.caption")',) * 2,
}

# The positional arguments of the three commands whose positional *tail* is variadic. Typst has no optional positional parameter, so `var`, `eq` and `val` take theirs through a `..rest` that no signature can describe -- the doc block names its slots instead.
#
# Only the positionals are written here: the named arguments are still read off the signature, so one added to the engine reaches both buttons with no edit, which is the whole point of deriving them.
POSITIONAL_TAILS = {
    "var": ('"R_k"', "7", '"kN"'),
    "eq": ('"R_d"', "`k_dens dot R_k`"),
    "val": ('"s_k"',),
}

# What an argument is written as where one command means by a name something another command means differently: var-unit's `key` is a name in the report's own registry and not a key of the glossary, and render-formula's `step` is a step of a calculation rather than a string.
ARGUMENT_LITERALS = {
    ("var-unit", "key"): '"R_k"',
    ("var-unit", "unit"): '"N"',
    ("unit", "written"): '"kN"',
    ("var-domain", "name"): '"S10 - 0.55 kN Ende"',
    ("verify", "relation"): "`F_Ed <= R_d`",
    ("fmt-number", "value"): "1.729",
    ("fmt-number", "unit"): '"kN/m²"',
    ("fmt-decimals", "x"): "1.7291232",
    ("fmt-decimals", "digits"): "3",
    ("render-formula", "step"): '(formula: "q_(k,11) = C_e dot C_T dot s_k")',
    ("render-prose", "sentence"): '"$c_\\"red\\"$ ist der Reduktionsfaktor"',
    ("t-pattern", "key"): '"connection.caption"',
    # distance's `connector` is a kind of fastener, not a connector of the export, and its d a diameter in mm (v0.87)
    ("distance", "d"): "8",
    ("distance", "connector"): '"nail"',
}

# Where a command reads one field of a var, the var sampled for it must carry that field: get-description and get-choices of a result (q_k11) printed nothing at all, the one var of the call carrying neither (v0.86). The sample is the first var of the instance that carries a description and choices.
SAMPLE_FOR = {
    ("get-description", "v"): "v_described",
    ("get-choices", "v"): "v_described",
    # `id` is too common a parameter name to sample globally -- it would fill every command that takes one.
    # `plan` names a position plan everywhere but here: a fire-protection plan is a plan of its own export (v0.115)
    ("show-fire-plan", "plan"): "fire_plan",
    # a member check saved in Woodcore (calc/domains/member_check.py)
    ("show-check-member", "id"): "check",
    # a steel check saved in Woodcore's «Acier» tabs (calc/domains/steel_check.py)
    ("show-check-steel", "id"): "steel_check",
}

# What an argument is written as where nothing in the export names it. A glossary key and a term of the data are values like any other; they simply live in calc/i18n/strings.json rather than in data/.
ARGUMENT_DEFAULTS = {
    "key": '"governing"',
    "lang": '"fr"',
    "text": '"Bezugshöhe"',
    "sentence": '"Massgebend"',
}


def sample_arguments(data_dir):
    """A real value for every argument that names something the export lists, so the call Woodcore copies is one that compiles (docs/ui/window.md §10quater).

    An argument the export cannot answer for -- a check, on a project that exports none (section 14) -- keeps a "<name>" placeholder, which is visibly not a value.
    """
    data = Path(data_dir)
    samples = {}

    def first(seq):
        return next(iter(seq), None)

    def load(fn, *args):
        try:
            return fn(*args)
        except (OSError, ValueError, KeyError):
            return None

    calculations = (load(load_json, data / "calculations.json") or {}).get("domains", {})
    domain = "snow-load" if "snow-load" in calculations else first(calculations)
    if domain:
        samples["domain"] = typst_str(domain)
        instances = calculations[domain].get("instances", {})
        instance = first(instances)
        if instance:
            samples["instance"] = typst_str(instance)
            record = instances[instance]
            name = first(record.get("results") or ()) or first(record.get("vars", {}))
            if name:
                samples["name"] = typst_str(name)
                var_of = f"var-of({typst_str(domain)}, {typst_str(instance)}, {typst_str(name)})"
                samples["v"] = samples["x"] = var_of
            described = first(k for k, v in record.get("vars", {}).items()
                               if v.get("description") and v.get("choices"))
            if described:
                samples["v_described"] = f"var-of({typst_str(domain)}, {typst_str(instance)}, {typst_str(described)})"
    for d in calculations.values():
        for record in d.get("instances", {}).values():
            for check in record.get("checks", ()):
                samples.setdefault("check-name", typst_str(check["name"]))
    live = calculations.get(LIVE_LOAD_DOMAIN, {}).get("instances", {})
    category = "A1" if "A1" in live else first(live)
    if category:
        samples["category"] = typst_str(category)

    records = list(load(load_positions, data) or ())
    positions = [p["position"] for p in records]
    if positions:
        samples["position"] = samples["positions"] = typst_str(positions[0])
    # the member checks (v0.88) read a timber position: the first one whose strength class Tabelle 8 or 9 carries, whose section the export read as a rectangle or a round one (v0.94) and which states a moisture class -- the one #member computes
    grades = set()
    for name in ("sia265-2021-tab08-vollholz", "sia265-2021-tab09-brettschichtholz"):
        tbl = load(load_json, ROOT / "calc" / "norms" / (name + ".json")) or {}
        for key, value in tbl.items():
            if key.startswith("table_"):
                grades |= {g["id"] for g in value.get("grades", ())}
    timber = first(p["position"] for p in records if p.get("strength_class") in grades and (p.get("section") or {}).get("type") in ("rec", "cir") and p.get("moisture_class"))
    if timber:
        samples["member"] = typst_str(timber)
    for key, loader, field in (
        ("connection", load_connections, "connections"),
        ("connector", load_connectors, "connectors"),
        ("plan", load_plans, "plans"),
        ("fire_plan", load_fire_plans, "plans"),
        ("model", load_statics, "models"),
        ("schema", load_schemas, "schemas"),
    ):
        entry = first((load(loader, data) or {}).get(field, {}))
        if entry:
            samples[key] = typst_str(entry)
    # A member check saved in Woodcore: the first one the export could calculate -- one it could not is refused by show-check-member, so a call naming it would not compile.
    checks = (load(load_json, data / "checks.json") or {}).get("checks", {})
    check = first(cid for cid, record in checks.items() if record.get("calculable"))
    if check:
        samples["check"] = typst_str(check)
    # A steel check saved in Woodcore (issues #50-#52): the first one the export could calculate, as for a member check.
    steel = (load(load_json, data / "steel.json") or {}).get("checks", {})
    steel_id = first(cid for cid, record in steel.items() if record.get("calculable"))
    if steel_id:
        samples["steel_check"] = typst_str(steel_id)
    # An excerpt of the annexes (v0.115): the first one the export found in its source -- one it could not find is refused by show-extract, so a call naming it would not compile.
    excerpts = (load(load_json, data / "annexes.json") or {}).get("excerpts", {})
    excerpt = first(eid for eid, record in excerpts.items() if record.get("resolved"))
    if excerpt:
        samples["excerpt"] = typst_str(excerpt)
    return samples


def argument_value(command, param, samples):
    """One argument, as the copied call writes it."""
    literal = ARGUMENT_LITERALS.get((command, param))
    if literal is not None:
        return literal
    special = SAMPLE_FOR.get((command, param))
    if special is not None and samples.get(special):
        return samples[special]
    return samples.get(param) or ARGUMENT_DEFAULTS.get(param) or typst_str(f"<{param}>")


# The member checks of SIA 265 4.2 (v0.88) name a member written down first, and take the actions a check cannot do without as named arguments defaulting to none -- Typst has no required named argument. The call copied writes the `#member` line before it and these values in both forms, so that it compiles as pasted.
MEMBER_CHECKS = {
    "verify-tension": {"N": "25"},
    "verify-compression": {"N": "-40", "l-k-y": "3000", "l-k-z": "3000"},
    "verify-compression-perp": {"F": "18", "l": "120"},
    "verify-bending": {"M-y": "4.5", "a": "3000"},
    "verify-bending-normal": {"N": "-40", "M-y": "5", "l-k-y": "3000", "l-k-z": "3000", "a": "3000"},
    "verify-shear": {"V": "12"},
    "verify-shear-off": {"F": "20", "l": "200"},
    "verify-shear-perp": {"V": "12", "sigma-90": "-1.2"},
}


def command_call(record, samples, complete):
    """One command's call: every argument with its default where `complete`, the required ones alone otherwise."""
    written = WRITTEN_CALLS.get(record["name"])
    if written is not None:
        return written[0 if complete else 1]
    name = record["name"]
    if name == "member" or name in MEMBER_CHECKS:
        return member_call(record, samples, complete)
    tail = POSITIONAL_TAILS.get(name)
    positional, named, body = [], [], False
    for param, default in split_params(record["params"] or ""):
        if param.startswith(VARIADIC):
            if tail is not None:
                positional += list(tail[len(positional):])  # the slots the doc block names
            continue  # otherwise variadic, and an empty one is what a call normally passes
        if default is not None:
            if complete:
                named.append(f"{param}: {default}")
        elif param == TRAILING_BODY:
            body = True
        elif tail is not None and len(positional) < len(tail):
            positional.append(tail[len(positional)])
        else:
            positional.append(argument_value(name, param, samples))
    call = f"#{name}({', '.join(positional + named)})"
    return call + "[\n  \n]" if body else call


def member_call(record, samples, complete):
    """A member check's call, or `#member`'s own, on the first timber position of the export."""
    name = record["name"]
    who = samples.get("member") or samples.get("position") or typst_str("<member>")
    required = MEMBER_CHECKS.get(name, {})
    named = []
    for param, default in split_params(record["params"] or "")[1:]:
        if param in required:
            named.append(f"{param}: {required[param]}")
        elif complete and default is not None:
            named.append(f"{param}: {default}")
    call = f"#{name}({', '.join([who] + named)})"
    return call if name == "member" else f"#member({who})" + chr(10) + call


def call_complete(record, samples):
    return command_call(record, samples, True)


def call_minimal(record, samples):
    return command_call(record, samples, False)


def import_line(here=None):
    """The import line every section starts with: `#import "../commands.typ": *` in a report alone, `#import "/commands.typ": *` in a project laid out in three documents, whose sections import the engine from the Typst root -- calc/documents.py's `import_line`, the one place it is spelt, for project `here`, else the current one, else none (where no project can be named, a report alone's line, which is COMMANDS_IMPORT_PATH's)."""
    calc = str(ROOT / "calc")
    if calc not in sys.path:
        sys.path.insert(0, calc)
    import documents
    import project
    if here is None:
        try:
            here = project.current()
        except project.ProjectError:
            here = None
    return documents.import_line(here)


def build_function_rows(records, samples):
    """Every top-level binding of the three lib files, the engine's own helpers included.

    An internal one carries its category, where it is and what it looks like, and nothing else: it is listed because `commands.typ`'s wildcard puts its name in every section's scope (docs/typst/command-docs.md §10sexies), not because anyone should call it.
    """
    rows = []
    for r in records:
        internal = r["category"] == INTERNAL
        block = parse_doc_block(r["doc_lines"])
        rows.append((
            r["category"], r["file"], r["line"], r["kind"], r["name"], r["signature"],
            "" if internal else block["summary"],
            "" if internal else block["returns"],
            "" if internal else call_minimal(r, samples),
            "" if internal else call_complete(r, samples),
        ))
    return rows


def build_parameter_rows(records):
    rows = []
    for r in records:
        if r["category"] == INTERNAL:
            continue
        defaults = dict(split_params(r["params"] or ""))
        for name, type_, meaning in parse_doc_block(r["doc_lines"])["bullets"]:
            rows.append((r["category"], r["name"], name, type_, defaults.get(name) or "", meaning.strip()))
    return rows


def build_example_rows(records):
    """One row per worked example, the code as its /// block carries it.

    `compiled` says whether Woodcore renders it and tests/test_commands.py holds it -- a ```typ fence -- or whether it is there to be read only (```typ-dry), which one command needs because nothing in this project exports a check for it to run against (section 14).
    """
    rows = []
    for r in records:
        if r["category"] == INTERNAL:
            continue
        for index, example in enumerate(parse_doc_block(r["doc_lines"])["examples"]):
            rows.append((r["category"], r["name"], index + 1,
                         "yes" if example["render"] else "no", example["code"]))
    return rows


def iter_sources(positions, calculations, connections=None):
    """Every {vars, steps, checks} record a show-* call can read, with the two arguments that select it: (domain, instance).

    Yields (source_type, domain, instance, engine, data) -- `sheet` for an instance read from a workbook, `norm` for one read from a norm table (a category of SIA 261's Tabelle 8, section 3.3septies), `position` for a domain carried by a position, `connection` for a connection saved in Woodcore (domain "connection", its id as the instance).
    """
    for domain, d in calculations["domains"].items():
        source_type = NORM_ENGINE if d["engine"] == NORM_ENGINE else "sheet"
        for sheet, instance in d["instances"].items():
            yield (source_type, domain, sheet, d["engine"], instance)
    for p in positions:
        for domain, data in p["domains"].items():
            yield ("position", domain, p["position"], "python", data)
    for cid, record in (connections or {}).get("connections", {}).items():
        yield ("connection", CONNECTION_DOMAIN, cid, "python", record)


def load_connections(data_dir):
    """data/connections.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "connections.json"
    return load_json(path) if path.exists() else {"connections": {}}


def load_connectors(data_dir):
    """data/connectors.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "connectors.json"
    return load_json(path) if path.exists() else {"connectors": {}}


def load_checks(data_dir):
    """data/checks.json, or none at all where it has not been exported yet."""
    path = Path(data_dir) / "checks.json"
    return load_json(path) if path.exists() else {"checks": {}}


def build_member_checks_rows(checks):
    """One row per member check saved in Woodcore: what it checks on which member, where each value came from, whether it can be calculated, and the call that prints it."""
    rows = []
    for cid, c in checks.get("checks", {}).items():
        rows.append([
            cid, c.get("name", ""), c.get("kind", ""), c.get("position") or "",
            c.get("grade") or "", c.get("grade_from") or "",
            (c.get("section") or {}).get("text", ""), c.get("section_from") or "",
            c.get("moisture_class"), c.get("moisture_from") or "",
            "%s %s" % ((c.get("N") or {}).get("value"), (c.get("N") or {}).get("unit")), c.get("alpha"),
            bool(c.get("calculable")), c.get("reason", ""),
            call_show_check_member(cid) if c.get("calculable") else "",
        ])
    return rows


def build_connectors_rows(connectors, connections):
    """One row per connector: what it is, **the detail schema it is drawn on**, the polar-moment calculations drawn on that same schema, and the call that prints it. A polar moment names no connector since v0.76, so the schema is what the two share."""
    def size(value):
        return value["value"] if value else None

    rows = []
    for cid, c in connectors.get("connectors", {}).items():
        schema = c.get("schema", "")
        calculations = [k for k, r in connections.get("connections", {}).items() if schema and r.get("schema") == schema]
        predrilled = {True: "Ja", False: "Nein"}.get(c["predrilled"], "")
        rows.append((
            cid, c["type"], c["assembly"], size(c["diameter"]), size(c["length"]), c["dimension"],
            c["steel"], predrilled, ", ".join(c["elements"]), c.get("schema", ""),
            ", ".join(m.get("element", "") for m in c.get("members", ())),
            ", ".join(f["id"] for f in c.get("fasteners", ())),
            ", ".join(calculations), call_show_connector(cid),
        ))
    return rows


def build_connections_rows(connections):
    """One row per connection: the schema it is drawn on, what it is, its polar moment, and the call that distributes forces over it (N, V, M to fill in)."""
    rows = []
    for cid, record in connections.get("connections", {}).items():
        vars_ = record["vars"]
        rows.append((
            cid, record.get("schema", ""), record.get("label", ""),
            vars_["n"]["value"], vars_["I_px"]["value"],
            vars_["I_pz"]["value"], vars_["I_p"]["value"], record["norm"], record["norm_ref"],
            call_show_polar_moment(cid),
        ))
    return rows


def build_calculations_rows(positions, calculations, connections=None):
    """One row per instance. `workbook` says which load workbook a sheet was read from -- several may be, so the sheet name alone does not -- and is empty for a position's domain and for one read from a norm table, as is `action` for a position's (permanent or variable, a property of a load domain, docs/domain/calculation-domains.md §3.3)."""
    def action(domain):
        return calculations["domains"].get(domain, {}).get("action", "")

    def example(domain, instance, d):
        # The Nutzlasten have a call of their own, which is how the report prints one (section 3.3septies); show-calc renders this domain too, and is what every other one is shown with -- and what a category the norm gives no values for (E) is shown with, that call being refused by the report.
        if domain == LIVE_LOAD_DOMAIN and "q_k" in d.get("vars", {}):
            return call_show_live_load(instance)
        return call_show_calc(domain, instance)

    return [
        (
            source_type, domain, instance, d.get("workbook", ""),
            action(domain), engine,
            len(d["vars"]), len(d["steps"]), len(d.get("results", [])),
            len(d["checks"]), len(d.get("tables", [])),
            example(domain, instance, d),
        )
        for source_type, domain, instance, engine, d in iter_sources(positions, calculations, connections)
    ]


def build_results_rows(positions, calculations, connections=None):
    """What show-calc prints at level 1, per instance."""
    return [
        (
            source_type, domain, instance, name,
            d["vars"][name]["symbol"], d["vars"][name]["value"],
            d["vars"][name]["unit"],
        )
        for source_type, domain, instance, _, d in iter_sources(positions, calculations, connections)
        for name in d.get("results", [])
    ]


def build_steps_rows(positions, calculations, connections=None):
    """The steps show-calc renders, in order. There is no example_call column: show-calc renders the whole instance, so a step is never addressed by name (docs/typst/api-conventions.md §10bis).
    """
    return [
        (
            source_type, domain, instance, i, s["name"], s["level"], s["result"],
            ", ".join(s["uses"]), s.get("formula", ""), s["norm"], s["norm_ref"],
        )
        for source_type, domain, instance, _, d in iter_sources(positions, calculations, connections)
        for i, s in enumerate(d["steps"], start=1)
    ]


def build_checks_rows(positions, calculations, connections=None):
    return [
        (
            source_type, domain, instance, c["name"], c["level"], c["norm"], c["norm_ref"],
            call_show_check(domain, instance, c["name"]),
        )
        for source_type, domain, instance, _, d in iter_sources(positions, calculations, connections)
        for c in d["checks"]
    ]


def build_vars_rows(positions, calculations, connections=None):
    return [
        (
            source_type, domain, instance, vn, v["label"], v["unit"], v["symbol"],
            v["level"], "description" in v, "choices" in v,
            call_show_input(domain, instance, vn),
        )
        for source_type, domain, instance, _, d in iter_sources(positions, calculations, connections)
        for vn, v in d["vars"].items()
    ]


def build_choices_rows(positions, calculations, connections=None):
    return [
        (source_type, domain, instance, vn, v["symbol"], c["value"], c["desc"])
        for source_type, domain, instance, _, d in iter_sources(positions, calculations, connections)
        for vn, v in d["vars"].items()
        for c in v.get("choices") or []
    ]


SHEET_SPECS = [
    ("Functions", ["category", "file", "line", "kind", "function", "signature", "summary",
                   "returns", "minimal_call", "complete_call"], "Functions"),
    ("Parameters", ["category", "function", "param", "type", "default", "meaning"], "Parameters"),
    ("Examples", ["category", "function", "example_no", "compiled", "source"], "Examples"),
    (
        "Calculations",
        ["source_type", "domain", "instance", "workbook", "action", "engine", "num_vars", "num_steps",
         "num_results", "num_checks", "num_tables", "example_call"],
        "Calculations",
    ),
    (
        "Results",
        ["source_type", "domain", "instance", "result_var", "symbol", "value", "unit"],
        "Results",
    ),
    (
        "Steps",
        ["source_type", "domain", "instance", "order", "step_name", "level",
         "result_var", "uses", "formula", "norm", "norm_ref"],
        "Steps",
    ),
    (
        "Checks",
        ["source_type", "domain", "instance", "check_name", "level", "norm",
         "norm_ref", "example_call"],
        "Checks",
    ),
    (
        "Vars",
        ["source_type", "domain", "instance", "var_name", "label", "unit", "symbol",
         "level", "has_description", "has_choices", "example_call"],
        "Vars",
    ),
    (
        "Choices",
        ["source_type", "domain", "instance", "var_name", "symbol", "choice_value",
         "choice_desc"],
        "Choices",
    ),
    (
        "Connectors",
        ["connector", "type", "assembly", "diameter [mm]", "length [mm]", "dimension", "steel",
         "predrilled", "elements", "schema", "member_elements", "markers", "polar_moments_on_schema",
         "example_call"],
        "Connectors",
    ),
    (
        "MemberChecks",
        ["check", "name", "kind", "position", "grade", "grade_from", "section", "section_from",
         "moisture_class", "moisture_from", "N", "alpha", "calculable", "reason", "example_call"],
        "MemberChecks",
    ),
    (
        "Connections",
        ["connection", "schema", "label", "n", "I_px [mm²]", "I_pz [mm²]", "I_p [mm²]",
         "norm", "norm_ref", "example_call"],
        "Connections",
    ),
    (
        "Plans",
        ["plan", "label", "floor", "background", "scale", "elements", "axes", "labels",
         "not_in_positions", "example_call"],
        "Plans",
    ),
    (
        "FirePlans",
        ["plan", "label", "frame", "background", "scale", "zones", "walls", "doors", "marks",
         "example_call"],
        "FirePlans",
    ),
    (
        "Statics",
        ["model", "label", "points", "members", "supports", "loads", "dimensions", "shown",
         "measured", "from_variables", "unresolved", "example_call"],
        "Statics",
    ),
    (
        "Schemas",
        ["schema", "label", "elements", "connectors", "aids", "dimensions",
         "element_ids", "connector_ids", "unresolved", "example_call"],
        "Schemas",
    ),
]


def build_reference_workbook(data_dir, out_path, root=ROOT):
    """Writes the catalogue of `root`'s lib/ and of the calls `data_dir`'s export can answer, to `out_path` -- the program's commands and the project's ids, which since v0.81 need not sit in one folder."""
    data_dir, out_path = Path(data_dir), Path(out_path)

    # All three files commands.typ re-exports, documented or not: until v0.64 this read calc-engine.typ alone and could not read a `#let` without parentheses, so var, eq, verify, val, unit, var-domain, var-unit, the ten commands of i18n.typ and the four settings of the document were in no catalog at all.
    records = lib_bindings(root)
    samples = sample_arguments(data_dir)
    positions = load_positions(data_dir)
    calculations = load_json(data_dir / "calculations.json")
    connections = load_connections(data_dir)
    connectors = load_connectors(data_dir)

    rows_by_sheet = {
        "Functions": build_function_rows(records, samples),
        "Parameters": build_parameter_rows(records),
        "Examples": build_example_rows(records),
        "Calculations": build_calculations_rows(positions, calculations, connections),
        "Results": build_results_rows(positions, calculations, connections),
        "Steps": build_steps_rows(positions, calculations, connections),
        "Checks": build_checks_rows(positions, calculations, connections),
        "Vars": build_vars_rows(positions, calculations, connections),
        "Choices": build_choices_rows(positions, calculations, connections),
        "Connectors": build_connectors_rows(connectors, connections),
        "MemberChecks": build_member_checks_rows(load_checks(data_dir)),
        "Connections": build_connections_rows(connections),
        "Plans": build_plans_rows(load_plans(data_dir)),
        "FirePlans": build_fire_plans_rows(load_fire_plans(data_dir)),
        "Statics": build_statics_rows(load_statics(data_dir)),
        "Schemas": build_schemas_rows(load_schemas(data_dir)),
    }

    # Built only where what it would hold changed since the last time it was written here: its rows are compared by their digest, kept beside the window's caches, with the file's own time -- so an export that changed nothing in the catalogue neither imports openpyxl (two seconds, with numpy) nor builds the workbook (docs/architecture/performance.md, v0.117).
    digest = _rows_digest(rows_by_sheet)
    if _unchanged(out_path, digest):
        return out_path
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, key in SHEET_SPECS:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for row in rows_by_sheet[key]:
            ws.append(row)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    if not _same_workbook(out_path, buffer.getvalue()):
        out_path.write_bytes(buffer.getvalue())
    _remember(out_path, digest)
    return out_path


DIGESTS_FILE = "reference-digests.json"


def _rows_digest(rows_by_sheet):
    import hashlib
    payload = json.dumps([SHEET_SPECS, rows_by_sheet], default=str, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _digests_path():
    """Where the digests are kept: the window's caches (%LOCALAPPDATA%\\Woodcore), never the project, which NextCloud syncs; None where that folder cannot be named."""
    try:
        try:
            import usersettings
        except ImportError:   # the export imports this module as tools.build_reference, calc/tools not on its path
            from tools import usersettings
        return usersettings.cache_dir(DIGESTS_FILE)
    except Exception:  # broad: no cache: the workbook is built every time, as before
        return None


def _stamp(out_path):
    try:
        stat = out_path.stat()
        return [stat.st_size, stat.st_mtime_ns]
    except OSError:
        return None


def _unchanged(out_path, digest):
    """Whether the workbook at `out_path` is the one written from rows of this `digest`, untouched since."""
    path = _digests_path()
    stamp = _stamp(out_path)
    if path is None or stamp is None:
        return False
    try:
        kept = json.loads(path.read_text(encoding="utf-8")).get(str(out_path))
    except (OSError, ValueError, AttributeError):
        return False
    return isinstance(kept, dict) and kept.get("digest") == digest and kept.get("stamp") == stamp


def _remember(out_path, digest):
    path = _digests_path()
    stamp = _stamp(out_path)
    if path is None or stamp is None:
        return
    try:
        kept = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
        if not isinstance(kept, dict):
            kept = {}
    except (OSError, ValueError):
        kept = {}
    kept[str(out_path)] = {"digest": digest, "stamp": stamp}
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(kept, indent=1), encoding="utf-8")
    except OSError:
        pass


# openpyxl stamps the time of every save into docProps/core.xml, so two builds of the same catalogue differ by that part alone. Written only when something else differs, an export leaves the committed workbook as it was, and a run of the test suite no longer leaves it modified (periodic audit 260916, AUD-1-3).
TIMESTAMPED_PART = "docProps/core.xml"


def _same_workbook(path, data):
    """True when the workbook at `path` holds exactly what `data` holds, its save time aside."""
    if not path.exists():
        return False
    try:
        with zipfile.ZipFile(path) as old, zipfile.ZipFile(io.BytesIO(data)) as new:
            names = set(old.namelist())
            if names != set(new.namelist()):
                return False
            return all(old.read(name) == new.read(name) for name in names if name != TIMESTAMPED_PART)
    except (OSError, zipfile.BadZipFile):
        return False


def main():
    # the project whose export is catalogued: --project <file>.woodcore, else the one project.current() finds
    sys.path.insert(0, str(ROOT / "calc"))
    import project
    named, _rest = project.take_argument(sys.argv[1:])
    here = project.use(named) if named else project.current()
    out_path = build_reference_workbook(here.data_dir, here.docs_dir / "typst-reference.xlsx")
    print(f"build_reference: wrote {out_path}")


if __name__ == "__main__":
    main()
